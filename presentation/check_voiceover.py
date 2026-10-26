"""Checks presentation/voiceover.md against the deck: one section per slide, same titles, 5-7 minutes.

    python presentation/check_voiceover.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
WPM = (140, 150)          # normal speaking pace
TARGET_MIN = (5.0, 7.0)


def sections(md: str) -> list[tuple[int, str, str]]:
    parts = re.split(r"^## Slide (\d+) — (.+)$", md, flags=re.M)
    return [(int(parts[i]), parts[i + 1].strip(), parts[i + 2].strip()) for i in range(1, len(parts), 3)]


def slide_titles(pptx: Path) -> list[str]:
    from pptx import Presentation  # python-pptx
    from pptx.enum.shapes import PP_PLACEHOLDER

    def title(slide) -> str:
        for sh in slide.placeholders:
            if sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                return sh.text_frame.text
        return ""

    return [title(sl) for sl in Presentation(str(pptx)).slides]


def main() -> int:
    secs = sections((HERE / "voiceover.md").read_text(encoding="utf-8"))
    titles = slide_titles(HERE / "setu_x.pptx")
    ok = True
    if [s[0] for s in secs] != list(range(1, len(titles) + 1)):
        print(f"FAIL: voiceover has sections {[s[0] for s in secs]}, deck has {len(titles)} slides")
        ok = False
    if len(titles) > 5:
        print(f"FAIL: deck has {len(titles)} slides (max 5)")
        ok = False
    total = 0
    for (n, title, body), deck_title in zip(secs, titles, strict=False):
        words = len(re.findall(r"[A-Za-z0-9][\w'.-]*", body))
        total += words
        same = title == deck_title
        ok &= same
        print(f"slide {n}: {words:4d} words · title {'matches' if same else f'DIFFERS (deck: {deck_title!r})'}")
    lo, hi = total / WPM[1], total / WPM[0]
    fits = TARGET_MIN[0] <= lo and hi <= TARGET_MIN[1]
    ok &= fits
    target = f"{TARGET_MIN[0]:.0f}–{TARGET_MIN[1]:.0f} min"
    print(f"total {total} words ≈ {lo:.1f}–{hi:.1f} min at {WPM[0]}–{WPM[1]} wpm (target {target}): "
          f"{'OK' if fits else 'OUT OF RANGE'}")
    print("PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
